import numpy as np
import SeatingPlan
import time
import math
from openpyxl import load_workbook
from numba.typed import List

wb = load_workbook('guest_seating_preferences_120.xlsx')
sh = wb.active

allVisitors = {}
visitorPrefs = {}

for col in sh.iter_cols(min_col=1, max_col=1, min_row=2):
    for cell in col:
        allVisitors[cell.value] = cell.row-1

lastKey = sh.max_row

for col in sh.iter_cols(min_col=2, max_col=2, min_row=2):
    for cell in col:
        friends = [friend.strip() for friend in cell.value.split(",")]
        for i in range(len(friends)):
            if(friends[i] not in allVisitors):
                allVisitors[friends[i]] = lastKey
                lastKey += 1

for col in sh.iter_cols(min_col=1, max_col=1, min_row=2):
    for cell in col:
        adjCell = sh.cell(row=cell.row, column=2)
        valGuest = cell.value
        valFriends = adjCell.value

        guestNum = allVisitors.get(valGuest)
        friendsNames = [friend.strip() for friend in valFriends.split(",")]
        friendsNum = tuple([allVisitors.get(friend) for friend in friendsNames])

        visitorPrefs[guestNum] = friendsNum

prefsMatrix = np.zeros((len(allVisitors), len(allVisitors)), dtype=int)

for i in range(1, len(visitorPrefs)+1):
    prefs = visitorPrefs.get(i)

    for col in prefs:
        prefsMatrix[i-1][col-1] = 1

startTime = time.perf_counter()

numStudents = len(allVisitors)
gridSize = math.ceil(math.sqrt(numStudents))

prefsList = List()
for value in visitorPrefs.values():
    arr = np.array(value, dtype=np.int64)
    prefsList.append(arr)

plan = SeatingPlan.SeatingPlan(numStudents,gridSize,gridSize,prefsMatrix,prefsList,1)

#Ask for: a copy of an old excel sheet that has all relevant information, ask for that specifically
#Ms Evans will change people's names
#Also need a map of table layout; how many people max per table
#Good to ask for actual seating plan
#Who should I ask for the excel sheet?
#What did people think about the layout last year? Any personal input?
#Talk to MR. N via email. 

#File IO to easily import large tables
#Optimize fitness

seating = np.array([
    [105,  82,  58, 103,  48,  39,  32,  72,  21,  67,  92],
    [ 89,  81,  22,  47, 114,  62,  23, 119,  73,  97,  61],
    [ 35, 101,  96,  64,  86,  34,  17, 113,  83, 111,  95],
    [110,  25,  78,  36,  70,  88,  42,  74, 116,  45,  13],
    [ 66,   8,  51,   9,   6,  10,  60,  76,  31,  59, 118],
    [ 99,  38,   5,  20,  14,  27, 106,  44,  79, 100,  52],
    [ 56,   3,  94,  68, 102, 108,  87,  71, 112,  46,  28],
    [ 91,   1,  29,  11,  63,  12,  55, 120,  85,  16,   7],
    [ 54, 117,  75,   4, 104,  37,  90, 107,  69,  49,  33],
    [ 15,   2,  77,  41,  43,  84,  40,  80,  18,  50,  57],
    [109,  24,  26,  98,  65, 115,  19,  53,  30,  93,   0]
])

##print(SeatingPlan.fitnessScore(seating, prefsMatrix))


numToVisitor = dict(zip(allVisitors.values(),allVisitors.keys()))

best = plan.annealRep()

endTime = time.perf_counter()

print("\nBest arrangement with people's names instead of numbers: ")
for i in range(best.shape[0]):
    print("In row " + str(i+1), end=": ")
    for elem in best[i]:
        if elem == 0:
            print("Empty", end=", ")
        else:
            print(numToVisitor[int(elem)], end=", ")
    print()
print("\nFound in " + str(endTime-startTime) + " seconds")     

